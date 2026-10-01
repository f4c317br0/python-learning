import openpyxl
from collections import OrderedDict

wb = openpyxl.load_workbook('census.xlsx')
ws = wb['Жители']

ages = []
heights = []
race_counts = OrderedDict()

for row in ws.iter_rows(min_row=2, values_only=True):
    if row[0] is None:
        continue
    _id, name, race, age, height = row[:5]
    ages.append(age)
    heights.append(height)
    race_counts[race] = race_counts.get(race, 0) + 1

max_age = max(ages)
avg_age = round(sum(ages) / len(ages), 2)
max_height = max(heights)
avg_height = round(sum(heights) / len(heights), 2)

title_race = None
best_count = -1
for race, count in race_counts.items():
    if count > best_count:
        best_count = count
        title_race = race

total = len(ages)
title_percent = round(best_count / total * 100, 2)

stats_ws = wb.create_sheet('Статистика')

rows = [
    ('Максимальный возраст', max_age),
    ('Средний возраст', avg_age),
    ('Максимальный рост', max_height),
    ('Средний рост', avg_height),
    ('Титульная раса (код)', title_race),
    ('Процент титульной расы', title_percent),
]

for i, (label, value) in enumerate(rows, start=1):
    stats_ws.cell(row=i, column=1, value=label)
    stats_ws.cell(row=i, column=2, value=value)

wb.save('statistics.xlsx')
