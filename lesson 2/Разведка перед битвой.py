import random
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill

enemies = input().split(', ')
qwe = {}
for name in enemies:
    qwe[name] = (random.randint(50, 500), random.randint(1, 10))
wb = Workbook()
ws = wb.active
ws['A1'] = 'Тип войск'
ws['B1'] = 'Количество'
ws['C1'] = 'Опасность'
for idx, name in enumerate(enemies):
    ws[f'A{idx + 2}'] = name
    ws[f'B{idx + 2}'] = qwe[name][0]
    ws[f'C{idx + 2}'] = qwe[name][1]
ws[f'A{len(enemies) + 2}'] = 'ИТОГО'
ws[f'B{len(enemies) + 2}'] = sum(qwe[el][0] for el in qwe)
wb.save('enemies.xlsx')
