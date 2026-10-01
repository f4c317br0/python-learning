import csv
import copy
import re
from docx import Document
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK

CSV_PATH = "books.csv"
TEMPLATE_PATH = "book_card.docx"
OUTPUT_PATH = "library.docx"


def fill_placeholders(doc, data):
    def fill_paragraph(paragraph):
        if not paragraph.runs:
            return
        text = "".join(run.text for run in paragraph.runs)

        def sub(match):
            return str(data.get(match.group(1), match.group(0)))

        new_text = re.sub(r"\{\{\s*(\w+)\s*\}\}", sub, text)

        if new_text != text:
            paragraph.runs[0].text = new_text
            for run in paragraph.runs[1:]:
                run.text = ""

    for p in doc.paragraphs:
        fill_paragraph(p)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    fill_paragraph(p)


def append_before(anchor, source_doc):
    for element in list(source_doc.element.body):
        if element.tag == qn("w:sectPr"):
            continue
        anchor.addprevious(copy.deepcopy(element))


def add_page_break(anchor):
    tmp_doc = Document()
    paragraph = tmp_doc.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    anchor.addprevious(copy.deepcopy(paragraph._p))


def load_books(csv_path):
    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f, delimiter=";"))

    books = []
    for row in rows[1:]:
        if len(row) < 5:
            continue
        book_id, title, author, year, copies = row[:5]
        if not copies.isdigit() or int(copies) <= 1:
            continue
        books.append({
            "ID": book_id,
            "Title": title,
            "Author": author,
            "Year": year,
            "Copies": copies,
        })
    return books


books = load_books(CSV_PATH)

result_doc = Document()
body = result_doc.element.body
sect_pr = body.find(qn("w:sectPr"))

for i, book in enumerate(books):
    card = Document(TEMPLATE_PATH)
    fill_placeholders(card, book)
    append_before(sect_pr, card)
    if i < len(books) - 1:
        add_page_break(sect_pr)

result_doc.save(OUTPUT_PATH)
